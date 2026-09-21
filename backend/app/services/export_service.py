"""Spreadsheet exports (ported from Round 1's export_service).

* the blank team-import template
* the team login sheet handed out after an import
* the coordinator-only results export (with fouls and power usage)
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Sequence

from app.core.exceptions import AppError
from app.services.results_service import format_elapsed

_HEADER_FILL = "FF111827"
_HEADER_FONT = "FFFFFFFF"
_HIGHLIGHT_FILL = "FFDCFCE7"


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise AppError("Excel export needs openpyxl: run `pip install openpyxl` and restart the API.") from exc
    return openpyxl, Alignment, Font, PatternFill, get_column_letter


def _write_sheet(
    sheet,
    *,
    title_lines: Sequence[str],
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    widths: Sequence[int],
    highlight_rows: set[int] | None = None,
) -> None:
    _mod, Alignment, Font, PatternFill, get_column_letter = _openpyxl()
    row_cursor = 1
    for index, line in enumerate(title_lines):
        cell = sheet.cell(row=row_cursor, column=1, value=line)
        cell.font = Font(bold=index == 0, size=14 if index == 0 else 10, color="FF334155")
        row_cursor += 1
    if title_lines:
        row_cursor += 1

    header_row = row_cursor
    fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color=_HEADER_FONT)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_cursor += 1

    highlight = PatternFill("solid", fgColor=_HIGHLIGHT_FILL)
    for body_index, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_cursor, column=col, value=value)
            if isinstance(value, str) and value.startswith("="):
                # openpyxl would store this as a live formula; a team named
                # =HYPERLINK(...) must stay plain text in a sheet of passwords.
                cell.data_type = "s"
            if highlight_rows and body_index in highlight_rows:
                cell.fill = highlight
        row_cursor += 1

    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if rows:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"


def _workbook_bytes(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_import_template() -> bytes:
    openpyxl, Alignment, Font, _fill, _letter = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Teams"
    _write_sheet(
        sheet,
        title_lines=[],
        headers=["Team Code", "Team Name", "Leader Name", "Phone Number", "Email"],
        rows=[
            ["", "Dragon Warriors", "Praveen Raja", "9876543210", "praveen@example.com"],
            ["", "Border Runners", "Asha Menon", "9123456780", "asha@example.com"],
        ],
        widths=[16, 28, 24, 18, 30],
    )
    help_sheet = workbook.create_sheet("How to use")
    help_sheet.column_dimensions["A"].width = 100
    lines = [
        "Filling in this sheet",
        "",
        "1. Delete the two sample rows on the 'Teams' sheet, then add one row per team.",
        "2. Column order does not matter - columns are matched by their header names,",
        "   so you can also upload Round 1's 'Round 2 Qualifiers' export as-is.",
        "3. Team Name and Phone Number are required. Team Code, Leader Name and Email are optional.",
        "",
        "What the server fills in for you",
        "",
        "  Team Code - kept if the sheet has one (so Round 1 teams keep their login),",
        "              otherwise generated in the form B@GCEE-1234#",
        "  Password  - the first 4 digits of the leader's phone number (same rule as Round 1)",
        "",
        "Nothing is created until you review the preview and confirm the import.",
    ]
    for row_index, line in enumerate(lines, start=1):
        cell = help_sheet.cell(row=row_index, column=1, value=line)
        cell.alignment = Alignment(wrap_text=False, vertical="top")
        cell.font = Font(bold=bool(line) and not line.startswith((" ", "1", "2", "3")), color="FF111827")
    workbook.active = 0
    return _workbook_bytes(workbook)


def build_credentials_workbook(*, teams: Sequence[dict], generated_at: datetime, title: str = "Team Logins") -> bytes:
    openpyxl, *_ = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Team Logins"
    _write_sheet(
        sheet,
        title_lines=[
            title,
            f"{len(teams)} team(s). Password = first 4 digits of the leader's phone number. "
            f"Generated {generated_at.strftime('%d %b %Y, %H:%M UTC')}.",
        ],
        headers=["Team Code", "Password", "Team Name", "Leader", "Phone", "Email"],
        rows=[
            [t.get("team_code"), t.get("password"), t.get("team_name"), t.get("leader_name"), t.get("leader_phone"), t.get("leader_email")]
            for t in teams
        ],
        widths=[16, 12, 28, 22, 16, 30],
    )
    return _workbook_bytes(workbook)


def build_results_workbook(*, event_name: str, rows: Sequence[dict], generated_at: datetime) -> bytes:
    openpyxl, *_ = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Results (coordinator)"
    finishers = {i for i, r in enumerate(rows) if r["finished"]}
    _write_sheet(
        sheet,
        title_lines=[
            f"{event_name} - full results (coordinator only: includes fouls)",
            f"Ranking: found the Joker, then fewest fouls, then fastest time. Generated {generated_at.strftime('%d %b %Y, %H:%M UTC')}.",
        ],
        headers=[
            "Rank", "Team Code", "Team Name", "Status", "Time", "Fouls", "Checkpoints",
            "Jack", "Queen", "King", "Attacks", "Defences", "Guides", "Photo hints",
            "Started (UTC)", "Finished (UTC)", "Leader", "Phone", "Email",
        ],
        rows=[
            [
                r["rank"], r["team_code"], r["team_name"], r["status"], format_elapsed(r["elapsed_s"]), r["foul_count"],
                f"{r['checkpoints']}/{r['total_checkpoints']}",
                "yes" if r["face_cards"]["JACK"] else "", "yes" if r["face_cards"]["QUEEN"] else "", "yes" if r["face_cards"]["KING"] else "",
                r["attacks_used"], r["defences_used"], r["guides_used"], r["photo_hints_used"],
                r["started_at"], r["completed_at"], r["leader_name"], r["leader_phone"], r["leader_email"],
            ]
            for r in rows
        ],
        widths=[6, 16, 26, 14, 10, 7, 12, 7, 7, 7, 9, 9, 7, 8, 24, 24, 20, 16, 28],
        highlight_rows=finishers,
    )
    return _workbook_bytes(workbook)
