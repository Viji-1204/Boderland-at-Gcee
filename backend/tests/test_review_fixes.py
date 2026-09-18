"""Regression tests for the independent code review's findings."""
from __future__ import annotations

import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.core.exceptions import RateLimitError, UnauthorizedError
from app.core.timeutil import utcnow
from app.database import SessionLocal
from app.models import Event, PowerUsage, RouteStop, Team
from app.services import auth_service, export_service, power_service
from app.services.team_import_service import ParsedSheet, build_rows, detect_columns
from tests.helpers import get_team, key, route_of, scan, seed, set_team

API = "/api/v1"


def test_rotating_client_address_cannot_reset_the_account_throttle(client):
    """Review #1: a spoofed X-Forwarded-For gave every guess a fresh throttle key."""
    seed(client)
    with SessionLocal() as db:
        for i in range(32):  # 8 per client x 4 = the per-account cap
            with pytest.raises(UnauthorizedError):
                auth_service.team_login(db, "B@GCEE-1003#", "0000", client=f"10.0.0.{i}")
        with pytest.raises(RateLimitError):
            auth_service.team_login(db, "B@GCEE-1003#", "9842", client="10.9.9.9")  # even the right password


def test_password_reset_logs_out_existing_sessions(client):
    """Review #10."""
    demo = seed(client)
    team = demo.teams["Dragon Warriors"]
    assert client.get(f"{API}/me/state", headers=team["headers"]).status_code == 200
    res = client.patch(f"{API}/admin/events/{demo.event_id}/teams/{team['id']}/password", json={"new_password": "5555"}, headers=demo.admin)
    assert res.status_code == 200
    stale = client.get(f"{API}/me/state", headers=team["headers"])
    assert stale.status_code == 401 and "password was changed" in stale.json()["detail"]
    token = team["headers"]["Authorization"].split()[1]
    from starlette.websockets import WebSocketDisconnect

    with client.websocket_connect(f"/ws/team?token={token}") as ws:
        with pytest.raises(WebSocketDisconnect):
            ws.receive_json()


def test_team_disqualified_before_lock_cannot_rejoin_after(client):
    """Review #2: it would start ACTIVE on a route the lock never validated."""
    demo = seed(client)
    team_id = demo.teams["Spade Squad"]["id"]
    base = f"{API}/admin/events/{demo.event_id}/teams/{team_id}"
    assert client.post(f"{base}/disqualify", json={"reason": "no-show"}, headers=demo.admin).status_code == 200
    with SessionLocal() as db:  # make its route invalid, as a stale DQ'd route would be
        db.query(RouteStop).filter(RouteStop.team_id == team_id).delete()
        db.commit()
    assert client.post(f"{API}/admin/events/{demo.event_id}/lock", headers=demo.admin).status_code == 200
    res = client.post(f"{base}/reinstate", headers=demo.admin)
    assert res.status_code == 409 and "before the configuration was locked" in res.json()["detail"]
    client.post(f"{API}/admin/events/{demo.event_id}/start", headers=demo.admin)
    assert get_team(team_id).status == "DISQUALIFIED"
    # and a routeless team never makes the radar crash
    set_team(team_id, status="ACTIVE", started_at=utcnow() - timedelta(seconds=5))
    radar = client.get(f"{API}/radar", params={"lat": 13.08, "lng": 80.27}, headers=demo.teams["Spade Squad"]["headers"])
    assert radar.status_code == 200 and radar.json()["locked"] is True


def test_start_only_activates_teams_validated_at_lock(client):
    demo = seed(client, lock=True)
    team_id = demo.teams["Heart Breakers"]["id"]
    set_team(team_id, status="NOT_STARTED")
    client.post(f"{API}/admin/events/{demo.event_id}/start", headers=demo.admin)
    assert get_team(team_id).status == "NOT_STARTED"
    assert get_team(demo.teams["Dragon Warriors"]["id"]).status == "ACTIVE"


def test_attack_without_power_reveals_nothing_about_the_target(client):
    """Review #8: rejections used to leak the target's freeze state first."""
    demo = seed(client, start=True)
    a, t = demo.teams["Joker's Wild"], demo.teams["Queen's Gambit"]
    set_team(t["id"], frozen_until=utcnow() + timedelta(minutes=5))
    with SessionLocal() as db:
        from app.models import TeamPower

        tp = db.query(TeamPower).filter(TeamPower.team_id == a["id"], TeamPower.kind == "ATTACK").one()
        tp.used = tp.owned
        db.commit()
    res = client.post(f"{API}/powers/attack", json={"target_team_id": t["id"], "idempotency_key": key()}, headers=a["headers"])
    assert res.status_code == 409 and res.json()["detail"] == "You have no Attack power left."


def test_joker_verified_during_pause_stops_the_clock_at_the_pause(client):
    """Review #6."""
    demo = seed(client, start=True)
    team = demo.teams["Dragon Warriors"]
    set_team(team["id"], status="FINAL", progress=8)
    client.post(f"{API}/admin/events/{demo.event_id}/pause", headers=demo.admin)
    with SessionLocal() as db:
        ev = db.get(Event, demo.event_id)
        ev.paused_at = utcnow() - timedelta(minutes=10)
        db.commit()
        paused_at = ev.paused_at
    client.post(f"{API}/admin/events/{demo.event_id}/teams/{team['id']}/verify-joker", headers=demo.admin)
    assert get_team(team["id"]).completed_at == paused_at


def test_coordinator_freeze_during_pause_lasts_exactly_as_long_as_set(client):
    demo = seed(client, start=True)
    team = demo.teams["Border Runners"]
    client.post(f"{API}/admin/events/{demo.event_id}/pause", headers=demo.admin)
    with SessionLocal() as db:
        ev = db.get(Event, demo.event_id)
        ev.paused_at = utcnow() - timedelta(minutes=10)
        db.commit()
    client.post(f"{API}/admin/events/{demo.event_id}/teams/{team['id']}/freeze", json={"duration_s": 60}, headers=demo.admin)
    client.post(f"{API}/admin/events/{demo.event_id}/resume", headers=demo.admin)
    left = (get_team(team["id"]).frozen_until - utcnow()).total_seconds()
    assert 55 <= left <= 61


def test_lost_race_on_an_expired_attack_releases_the_write_lock(client, monkeypatch):
    """Review #4: a zero-row UPDATE kept SQLite's write lock for the whole request."""
    demo = seed(client, start=True)
    a, t = demo.teams["Heart Breakers"], demo.teams["Spade Squad"]
    usage_id = client.post(f"{API}/powers/attack", json={"target_team_id": t["id"], "idempotency_key": key()}, headers=a["headers"]).json()["usage_id"]
    with SessionLocal() as db:
        db.get(PowerUsage, usage_id).expires_at = utcnow() - timedelta(seconds=1)
        db.commit()
    monkeypatch.setattr(power_service, "_resolve", lambda *args, **kwargs: False)  # someone else won
    with SessionLocal() as db:
        assert power_service.expire_due_attacks(db, db.get(Event, demo.event_id), utcnow()) == 0
        assert not db.in_transaction()


def test_formulas_in_exports_stay_plain_text():
    """Review #7: openpyxl would store '=HYPERLINK(...)' as a live formula."""
    from datetime import datetime, timezone

    content = export_service.build_credentials_workbook(
        teams=[{"team_code": "B@GCEE-0001#", "password": "1234", "team_name": '=HYPERLINK("http://evil","x")'}],
        generated_at=datetime.now(timezone.utc),
    )
    sheet = load_workbook(io.BytesIO(content)).active
    cells = [c for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


def test_import_flags_names_the_database_cannot_hold():
    """Review #11: Postgres rejects >120 characters and the whole import failed."""
    parsed = ParsedSheet(headers=["Team Name", "Phone"], rows=[["X" * 121, "9876543210"], ["Fine", "9123456789"]], header_row_number=1)
    rows = build_rows(parsed, detect_columns(parsed.headers), existing_names=[], existing_codes=[], capacity=60)
    assert [r.status for r in rows] == ["NAME_TOO_LONG", "OK"]


def test_development_signing_key_is_random_not_the_placeholder(tmp_path, monkeypatch):
    """Review #12: 'change-me' let anyone forge admin tokens on a dev server."""
    from app import config

    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    first = config._local_dev_secret()
    assert len(first) >= 32 and first not in config._PLACEHOLDER_SECRETS
    assert config._local_dev_secret() == first  # persisted, so logins survive restarts
