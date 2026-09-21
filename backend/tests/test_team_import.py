"""Round 1's spreadsheet import, ported - including Round 1's own qualifier export."""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from app.services import export_service
from app.services.team_import_service import (
    ParsedSheet,
    build_rows,
    detect_columns,
    generate_team_code,
    normalise_phone,
    parse_sheet,
    password_from_phone,
)
from tests.helpers import seed

API = "/api/v1"


def test_phone_rules_match_round1():
    assert normalise_phone("+91 98765 43210") == "9876543210"
    assert password_from_phone("98765 43210") == "9876"
    assert password_from_phone(9876543210.0) == "9876"
    assert password_from_phone("12") == ""


def test_generated_codes_are_round1_format():
    taken: set[str] = set()
    codes = {generate_team_code(taken) for _ in range(50)}
    assert len(codes) == 50 and all(c.startswith("B@GCEE-") and c.endswith("#") and len(c) == 12 for c in codes)


def test_column_detection_handles_round1_qualifier_export():
    headers = ["Seed", "Team Code", "Team Name", "Room", "Room Rank", "Leader", "Phone", "Email", "MindMaze", "Total"]
    mapping = detect_columns(headers)
    assert mapping["team_code"] == 1 and mapping["team_name"] == 2
    assert mapping["leader_name"] == 5 and mapping["leader_phone"] == 6 and mapping["leader_email"] == 7


def test_existing_codes_are_kept_and_duplicates_flagged():
    parsed = ParsedSheet(
        headers=["Team Code", "Team Name", "Phone"],
        rows=[["B@GCEE-7391#", "Alpha", "9876543210"], ["B@GCEE-7391#", "Beta", "9123456780"], ["", "Gamma", "9000000000"]],
        header_row_number=1,
    )
    rows = build_rows(parsed, detect_columns(parsed.headers), existing_names=[], existing_codes=[], capacity=60)
    assert rows[0].status == "OK" and rows[0].team_code == "B@GCEE-7391#" and rows[0].password == "9876"
    assert rows[1].status == "DUPLICATE_CODE"
    assert rows[2].status == "OK" and rows[2].team_code.startswith("B@GCEE-")


def _qualifier_sheet() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Round 2 Qualifiers"
    ws.append(["Round 1 - qualifiers"])  # title banner, like Round 1's export
    ws.append([])
    ws.append(["Seed", "Team Code", "Team Name", "Room", "Leader", "Phone", "Email", "Total"])
    ws.append([1, "B@GCEE-4242#", "Night Owls", "R01", "Priya", "+91 99887 76655", "p@x.com", 50])
    ws.append([2, "B@GCEE-5151#", "Sky Walkers", "R02", "Vijay", "9012345678", "", 48])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_import_round1_qualifiers_keeps_their_logins(client):
    demo = seed(client)
    files = {"file": ("qualifiers.xlsx", _qualifier_sheet(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post(f"{API}/admin/events/{demo.event_id}/teams/import/preview", files=files, headers=demo.admin).json()
    assert preview["header_row_number"] == 3 and preview["importable_count"] == 2
    rows = [r for r in preview["rows"] if r["status"] == "OK"]
    assert [r["team_code"] for r in rows] == ["B@GCEE-4242#", "B@GCEE-5151#"]

    created = client.post(f"{API}/admin/events/{demo.event_id}/teams/import", json={"rows": rows}, headers=demo.admin)
    assert created.status_code == 201 and created.json()["created_count"] == 2
    login = client.post(f"{API}/auth/team/login", json={"team_code": "B@GCEE-4242#", "password": "9988"})
    assert login.status_code == 200 and login.json()["team_name"] == "Night Owls"

    sheet = client.post(f"{API}/admin/teams/credentials-export", json=created.json(), headers=demo.admin)
    wb = load_workbook(io.BytesIO(sheet.content))
    values = [c for row in wb.active.iter_rows(values_only=True) for c in row]
    assert "B@GCEE-4242#" in values and "9988" in values


def test_template_round_trips_through_the_importer():
    parsed = parse_sheet(export_service.build_import_template(), "template.xlsx")
    mapping = detect_columns(parsed.headers)
    assert {"team_code", "team_name", "leader_name", "leader_phone", "leader_email"} <= set(mapping)
    assert len(parsed.rows) == 2


def test_teams_are_dealt_different_secret_sentences(client):
    """Nobody has to type a sentence: one is dealt per team, no two alike,
    from the Alice in Borderland pool - unless the coordinator types one."""
    from app.services import sentences
    from app.services.event_service import split_sentence

    demo = seed(client)
    eid = demo.event_id
    base = f"{API}/admin/events/{eid}/teams"
    for t in client.get(base, headers=demo.admin).json():  # start from an empty game
        client.delete(f"{base}/{t['id']}", headers=demo.admin)

    assert len(sentences.SENTENCES) == 20
    for line in sentences.SENTENCES:
        for parts in (7, 8, 9):
            assert split_sentence(line, parts) is not None, line  # every line splits for 7-9 checkpoints

    dealt = []
    for i in range(20):
        res = client.post(base, json={"team_name": f"Team {i:02d}", "password": "1234"}, headers=demo.admin)
        assert res.status_code == 201, res.text
        dealt.append(res.json()["sentence"])
    assert set(dealt) == set(sentences.SENTENCES)  # twenty teams, twenty different sentences
    # The twenty-first reuses one; a typed sentence is kept as typed.
    extra = client.post(base, json={"team_name": "Team 21", "password": "1234"}, headers=demo.admin).json()
    assert extra["sentence"] in sentences.SENTENCES
    typed = client.post(base, json={"team_name": "Team 22", "password": "1234", "sentence": "MY OWN WORDS ONE TWO THREE FOUR FIVE SIX SEVEN"}, headers=demo.admin).json()
    assert typed["sentence"] == "MY OWN WORDS ONE TWO THREE FOUR FIVE SIX SEVEN"


def test_import_without_a_sentence_column_deals_sentences_too(client):
    from app.services import sentences

    demo = seed(client)
    eid = demo.event_id
    csv_text = "Team Name,Leader Name,Leader Phone\nAlpha Hunters,Arun,9840012345\nBeta Runners,Divya,9841012345\n"
    files = {"file": ("teams.csv", csv_text.encode(), "text/csv")}
    preview = client.post(f"{API}/admin/events/{eid}/teams/import/preview", files=files, headers=demo.admin)
    assert preview.status_code == 200, preview.text
    rows = [r for r in preview.json()["rows"] if r["status"] == "OK"]
    assert len(rows) == 2
    commit = client.post(f"{API}/admin/events/{eid}/teams/import", json={"rows": rows}, headers=demo.admin)
    assert commit.status_code == 201, commit.text
    teams = {t["team_name"]: t for t in client.get(f"{API}/admin/events/{eid}/teams", headers=demo.admin).json()}
    got = [teams["Alpha Hunters"]["sentence"], teams["Beta Runners"]["sentence"]]
    assert all(s in sentences.SENTENCES for s in got) and got[0] != got[1]


def test_lock_deals_a_sentence_to_anyone_still_without_one(client):
    from app.services import sentences

    demo = seed(client)
    eid = demo.event_id
    team_id = demo.teams["Spade Squad"]["id"]
    client.patch(f"{API}/admin/events/{eid}/teams/{team_id}", json={"sentence": ""}, headers=demo.admin)
    checks = {c["key"]: c for c in client.get(f"{API}/admin/events/{eid}/readiness", headers=demo.admin).json()}
    assert checks["sentences"]["ok"] is True  # a missing one isn't a blocker: it is dealt at lock
    dealt = client.post(f"{API}/admin/events/{eid}/teams/deal-sentences", headers=demo.admin).json()
    assert dealt["dealt"] == 1  # the button on the Teams page
    assert client.post(f"{API}/admin/events/{eid}/teams/deal-sentences", headers=demo.admin).json()["dealt"] == 0
    client.patch(f"{API}/admin/events/{eid}/teams/{team_id}", json={"sentence": ""}, headers=demo.admin)
    assert client.post(f"{API}/admin/events/{eid}/lock", headers=demo.admin).status_code == 200  # lock deals it too
    state = client.get(f"{API}/me/state", headers=demo.teams["Spade Squad"]["headers"]).json()
    assert state["fragments_total"] == 8
    with_sentence = next(t for t in client.get(f"{API}/admin/events/{eid}/teams", headers=demo.admin).json() if t["id"] == team_id)
    assert with_sentence["sentence"] in sentences.SENTENCES
